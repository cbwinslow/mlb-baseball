"""Research Database Exporter & Interoperability Layer (EXPORT-01, Plan v1).

Provides high-performance, streaming multi-format exports for research data:
1. Validated relation allow-list (raw.*, core.*, gold.*) with zero arbitrary SQL.
2. Read-only repeatable-read streaming transactions via server-side cursors.
3. Multi-format serialisation: CSV, Excel (.xlsx with 1M-row safety guard), and Parquet.
4. Rights-filtered bundle export (`--profile public_safe`) with manifest and optional zip.
5. Defensive health checks wired into `mlb doctor`.
"""

from __future__ import annotations

import csv
import dataclasses
import json
import logging
import shutil
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import psycopg
from psycopg import sql

from mlb_baseball.db import fetch_one
from mlb_baseball.health import Check

logger = logging.getLogger(__name__)

RETROSHEET_ATTRIBUTION = (
    "Information used here was obtained free of charge from and is copyrighted by "
    "Retrosheet. Interested parties may contact Retrosheet at 20 Sunset Rd., Newark, DE 19711."
)

MAX_EXCEL_DATA_ROWS = 1_048_575  # 1,048,576 total rows minus 1 header row


@dataclasses.dataclass(frozen=True)
class ExportRelation:
    """An allow-listed database relation available for researcher export."""

    schema: str
    table: str
    season_column: str | None = None
    profile: str = "local_research"
    rights_note: str = ""
    primary_key: tuple[str, ...] | None = None

    @property
    def qualified_name(self) -> str:
        return f"{self.schema}.{self.table}"


# Canonical list of allowable export relations across raw, core, and gold schemas.
#
# profile="public_safe" is Retrosheet-only, per docs/SOURCE_RIGHTS.md: "public_safe
# currently permits only Retrosheet connector families." A relation is public_safe
# ONLY if every source that feeds it is Retrosheet. Anything touching Statcast, the
# MLB Stats API, Baseball-Reference, Lahman, Chadwick, or market data is
# local_research even when it also contains Retrosheet-derived columns -- core.game
# carries MLB-API weather, core.pitch is Statcast, core.player is conformed through
# the Chadwick register, gold.player_season is Baseball-Reference, and so on.
RELATIONS: tuple[ExportRelation, ...] = (
    # --- public_safe: Retrosheet-derived only (redistributable with attribution) ---
    # State->value lookup matrices, each built purely FROM raw.retrosheet_event
    # (+ raw.retrosheet_gameinfo for the game filter); see mlb_baseball/sql/*_matrix_build.sql.
    ExportRelation(
        "gold",
        "run_expectancy_24",
        None,
        "public_safe",
        "24 base-out-state run expectancy matrix (from Retrosheet events)",
    ),
    ExportRelation(
        "gold",
        "win_expectancy",
        None,
        "public_safe",
        "Win expectancy by score/inning/base-out state (from Retrosheet events)",
    ),
    ExportRelation(
        "gold",
        "leverage_index",
        None,
        "public_safe",
        "Leverage index matrix (from Retrosheet events)",
    ),
    ExportRelation(
        "raw", "retrosheet_event", None, "public_safe", "Raw Retrosheet play-by-play events"
    ),
    ExportRelation(
        "raw", "retrosheet_gameinfo", None, "public_safe", "Raw Retrosheet per-game metadata"
    ),
    # --- local_research: conformed dims/facts that mix in non-Retrosheet sources ---
    ExportRelation("gold", "game_export", "season", "local_research", "Wide per-game export view"),
    ExportRelation(
        "gold",
        "player_season",
        "season",
        "local_research",
        "Player season lines (Baseball-Reference + WAR)",
    ),
    ExportRelation(
        "gold", "team_season", "season", "local_research", "Team season stats (Lahman + Retrosheet)"
    ),
    ExportRelation(
        "gold", "division_standing", "season", "local_research", "Season-end division standings"
    ),
    # Grain-complete statistic backbone (Plan 03B). local_research, not
    # public_safe: the stat content is 100% Retrosheet events, but the
    # builders join the conformed core.game / core.player / core.team dims
    # for surrogate keys, and those dims mix in non-Retrosheet sources (the
    # public_safe bar is "every source that feeds it is Retrosheet"). A
    # public_safe variant keyed by retro ids is possible follow-up work.
    ExportRelation(
        "gold",
        "batting_game",
        "season",
        "local_research",
        "Batting box line per (game, player, team)",
        ("game_id", "player_id", "team_id"),
    ),
    ExportRelation(
        "gold",
        "pitching_game",
        "season",
        "local_research",
        "Pitching box line per (game, charged pitcher, team)",
        ("game_id", "player_id", "team_id"),
    ),
    ExportRelation(
        "gold",
        "batting_season",
        "season",
        "local_research",
        "Season batting line per (player, season, team) + combined row",
        ("id",),
    ),
    ExportRelation(
        "gold",
        "batting_team",
        "season",
        "local_research",
        "Season batting line per (team, season)",
        ("team_id", "season"),
    ),
    ExportRelation(
        "gold",
        "pitching_season",
        "season",
        "local_research",
        "Season pitching line per (player, season, team) + combined row",
        ("id",),
    ),
    ExportRelation(
        "gold",
        "pitching_team",
        "season",
        "local_research",
        "Season pitching line per (team, season)",
        ("team_id", "season"),
    ),
    ExportRelation(
        "gold",
        "batting_career",
        None,
        "local_research",
        "Career batting line per player",
        ("player_id",),
    ),
    ExportRelation(
        "gold",
        "pitching_career",
        None,
        "local_research",
        "Career pitching line per player",
        ("player_id",),
    ),
    ExportRelation(
        "core", "game", "season", "local_research", "Conformed games (Retrosheet + MLB API)"
    ),
    ExportRelation(
        "core", "play", "season", "local_research", "Conformed plays (Retrosheet + MLB API PBP)"
    ),
    ExportRelation("core", "pitch", "season", "local_research", "Conformed pitches (Statcast)"),
    ExportRelation(
        "core", "player", None, "local_research", "Player directory (conformed via Chadwick)"
    ),
    ExportRelation("core", "team", None, "local_research", "Franchise registry"),
    ExportRelation("core", "team_alias", None, "local_research", "Team alias crosswalk"),
    ExportRelation("core", "venue", None, "local_research", "Ballpark reference registry"),
    ExportRelation("core", "standing", "season", "local_research", "Conformed standings"),
    # Local Research Gold Relations (Excluded from public_safe)
    ExportRelation(
        "gold",
        "game_feature",
        "season",
        "local_research",
        "Enriched feature matrix (contains Statcast/MLB API data)",
    ),
    ExportRelation(
        "gold",
        "game_feature_snapshot",
        "season",
        "local_research",
        "Point-in-time feature snapshot",
    ),
    ExportRelation(
        "gold", "prediction", "season", "local_research", "Model game-winner forecast outputs"
    ),
    ExportRelation(
        "gold", "total_prediction", "season", "local_research", "Model run-total forecast outputs"
    ),
    # Local Research Core Relations
    ExportRelation(
        "core", "player_war", "season", "local_research", "Baseball-Reference WAR calculations"
    ),
    ExportRelation(
        "core", "market", "season", "local_research", "Prediction market comparison lines"
    ),
    # Local Research Raw Relations (MLB API, Chadwick, RSS)
    ExportRelation(
        "raw", "mlb_playbyplay", "_season", "local_research", "MLB Stats API play-by-play"
    ),
    ExportRelation(
        "raw", "mlb_boxscore_batting", "_season", "local_research", "MLB Stats API boxscore batting"
    ),
    ExportRelation(
        "raw",
        "mlb_boxscore_fielding",
        "_season",
        "local_research",
        "MLB Stats API boxscore fielding",
    ),
    ExportRelation("raw", "mlb_roster", "_season", "local_research", "MLB Stats API team rosters"),
    ExportRelation(
        "raw", "mlb_player_stat", "_season", "local_research", "MLB Stats API player season totals"
    ),
    ExportRelation(
        "raw", "mlb_team_stat", "_season", "local_research", "MLB Stats API team season totals"
    ),
    ExportRelation("raw", "mlb_standing", "_season", "local_research", "MLB Stats API standings"),
    ExportRelation(
        "raw", "mlb_game_context", "_season", "local_research", "MLB Stats API game context"
    ),
    ExportRelation(
        "raw", "mlb_game_pace", "_season", "local_research", "MLB Stats API game pace metrics"
    ),
    ExportRelation(
        "raw", "mlb_linescore", "_season", "local_research", "MLB Stats API linescore innings"
    ),
    ExportRelation(
        "raw", "mlb_stat_leader", "_season", "local_research", "MLB Stats API category leaders"
    ),
    ExportRelation(
        "raw", "mlb_team_leader", "_season", "local_research", "MLB Stats API team leaders"
    ),
    ExportRelation(
        "raw", "mlb_umpire", "_season", "local_research", "MLB Stats API umpire directory"
    ),
    ExportRelation(
        "raw", "mlb_win_prob", "_season", "local_research", "MLB Stats API win probability"
    ),
    ExportRelation("raw", "mlb_alumni", "_season", "local_research", "MLB Stats API alumni lists"),
    ExportRelation("raw", "mlb_coach", "_season", "local_research", "MLB Stats API coaching staff"),
    ExportRelation(
        "raw", "mlb_free_agent", "_season", "local_research", "MLB Stats API free agents"
    ),
    ExportRelation(
        "raw", "mlb_transaction", "_season", "local_research", "MLB Stats API transaction log"
    ),
    ExportRelation(
        "raw", "register_people", None, "local_research", "Chadwick Register player crosswalk"
    ),
    ExportRelation(
        "raw", "register_names", None, "local_research", "Chadwick Register name variants"
    ),
    ExportRelation(
        "raw", "register_links", None, "local_research", "Chadwick Register identifier links"
    ),
    ExportRelation(
        "raw", "register_countries", None, "local_research", "Chadwick Register country codes"
    ),
    ExportRelation(
        "raw", "news", None, "local_research", "MLB / Trade Rumors / ESPN RSS news feeds"
    ),
    ExportRelation("meta", "ingestion_run", "season", "local_research", "Ingestion run ledger"),
    ExportRelation("meta", "ingestion_item", "season", "local_research", "Ingestion item ledger"),
)

ALLOWLIST: dict[str, ExportRelation] = {r.qualified_name: r for r in RELATIONS}

# Grain-complete statistic backbone (Plan 03B) HF publish preset
# (delivery-surface change). Publish eligibility is a redistribution rights
# decision, reviewed in openspec/changes/delivery-surface/rights-review.md --
# distinct from ExportRelation.profile above, which governs per-source
# *intake* gating (MLB_DATA_PROFILE) and is a stricter, source-family-based
# test unsuited to a derived table's actual redistributable content.
BACKBONE_SCHEMA_VERSION = "1.0.0"

BACKBONE_CANDIDATES: tuple[str, ...] = (
    "gold.batting_game",
    "gold.pitching_game",
    "gold.batting_season",
    "gold.pitching_season",
    "gold.batting_team",
    "gold.pitching_team",
    "gold.batting_career",
    "gold.pitching_career",
    "gold.player_season",
    "gold.team_season",
)

# Excluded, with the reason recorded in the export manifest and log --
# see rights-review.md for the full evidence.
BACKBONE_EXCLUDED: dict[str, str] = {
    "gold.player_season": (
        "Sourced from raw.bref_batting/raw.bref_pitching (Baseball-Reference via "
        "pybaseball) and core.player_war (Baseball-Reference WAR). "
        "docs/SOURCE_RIGHTS.md records no redistribution permission for "
        "Baseball-Reference data."
    ),
    "gold.team_season": (
        "Traditional totals sourced from raw.lahman_teams; war sourced from "
        "core.player_war (Baseball-Reference). docs/SOURCE_RIGHTS.md records no "
        "automated public-safe approval for Lahman, and no permission for "
        "Baseball-Reference redistribution."
    ),
}

BACKBONE_TABLES: tuple[str, ...] = tuple(
    name for name in BACKBONE_CANDIDATES if name not in BACKBONE_EXCLUDED
)


def resolve_relation(name: str) -> ExportRelation:
    """Validate and resolve a relation name against the allow-list.

    Accepts fully qualified ("gold.game_export") or bare ("game_export") names.
    Raises ValueError on unknown or ambiguous relation names.
    """
    clean_name = name.strip()
    if clean_name in ALLOWLIST:
        return ALLOWLIST[clean_name]

    # Check for bare table name match
    matches = [r for r in RELATIONS if r.table == clean_name]
    if len(matches) == 1:
        return matches[0]
    elif len(matches) > 1:
        options = ", ".join(m.qualified_name for m in matches)
        raise ValueError(f"Ambiguous relation {name!r}. Specify schema: {options}")

    allowed = sorted(ALLOWLIST.keys())
    raise ValueError(
        f"Relation {name!r} is not in the export allow-list. "
        f"Allowed relations: {', '.join(allowed)}"
    )


def _build_select_query(
    rel: ExportRelation,
    season: int | None = None,
    order_by: tuple[str, ...] | None = None,
) -> tuple[sql.Composed, list[Any]]:
    """Construct a strictly-parameterized SQL query for the given relation.

    ``order_by`` (when given) must be a tuple of the relation's own column
    names -- callers pass ``rel.primary_key`` to get deterministic row order.
    """
    query = sql.SQL("SELECT * FROM {}.{}").format(
        sql.Identifier(rel.schema),
        sql.Identifier(rel.table),
    )
    params: list[Any] = []
    if season is not None:
        if rel.season_column is None:
            raise ValueError(
                f"Relation {rel.qualified_name!r} does not have a season column "
                "and cannot be filtered by --season"
            )
        query = sql.SQL("SELECT * FROM {}.{} WHERE {} = %s").format(
            sql.Identifier(rel.schema),
            sql.Identifier(rel.table),
            sql.Identifier(rel.season_column),
        )
        params.append(season)
    if order_by:
        query = sql.SQL("{} ORDER BY {}").format(
            query,
            sql.SQL(", ").join(sql.Identifier(col) for col in order_by),
        )
    return query, params


def _build_count_query(
    rel: ExportRelation, season: int | None = None
) -> tuple[sql.Composed, list[Any]]:
    """Construct a parameterized count query for the given relation."""
    query = sql.SQL("SELECT count(*) FROM {}.{}").format(
        sql.Identifier(rel.schema),
        sql.Identifier(rel.table),
    )
    params: list[Any] = []
    if season is not None and rel.season_column is not None:
        query = sql.SQL("SELECT count(*) FROM {}.{} WHERE {} = %s").format(
            sql.Identifier(rel.schema),
            sql.Identifier(rel.table),
            sql.Identifier(rel.season_column),
        )
        params.append(season)
    return query, params


def _set_read_only_repeatable_read(conn: psycopg.Connection) -> None:
    """Ensure the connection runs in a read-only, repeatable-read transaction."""
    try:
        with conn.cursor() as cur:
            cur.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
    except psycopg.errors.ActiveSqlTransaction:
        conn.commit()
        with conn.cursor() as cur:
            cur.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")


# PostgreSQL type OID -> pyarrow type, for the columns most common in this
# project's relations. Used only to correct a column pyarrow would otherwise
# infer as `null` (every value in the exported result is None) -- pyarrow's
# own per-batch type inference is left untouched for every column that has
# at least one non-null value. An untyped `null` column round-trips through
# duckdb as an unrelated type (e.g. INTEGER), a real schema-drift bug caught
# by the backbone export's pyarrow/duckdb round-trip test.
_PG_OID_TO_ARROW_TYPE: dict[int, str] = {
    16: "bool_",  # boolean
    20: "int64",  # bigint
    21: "int16",  # smallint
    23: "int32",  # integer
    700: "float32",  # real
    701: "float64",  # double precision
    # SHORTCUT: numeric -> float64, not decimal128(precision, scale). This
    # project's numeric columns (avg/obp/slg/era/... in the backbone tables)
    # are unconstrained `numeric` in the DDL (no declared precision/scale),
    # so psycopg's Cursor.description reports none to build an exact
    # decimal128 from; float64 has ample precision for a handful of rate-stat
    # decimal digits. Revisit if a numeric column needing exact/high-precision
    # decimal semantics is added to the backbone preset.
    1700: "float64",  # numeric
    1082: "date32",  # date
    1114: "timestamp_us",  # timestamp without time zone
    1184: "timestamp_us_utc",  # timestamptz
    25: "string",  # text
    1043: "string",  # varchar
    18: "string",  # "char"
    19: "string",  # name
}


def _arrow_type_for_oid(oid: int) -> Any:
    """Resolve a PostgreSQL column type OID to a concrete pyarrow type.

    Falls back to ``pa.string()`` for any OID not in the common-type table
    above -- safe because this is only ever applied to a column pyarrow
    already inferred as untyped ``null`` (no real values to lose fidelity on).
    """
    import pyarrow as pa

    name = _PG_OID_TO_ARROW_TYPE.get(oid, "string")
    if name == "bool_":
        return pa.bool_()
    if name == "timestamp_us":
        return pa.timestamp("us")
    if name == "timestamp_us_utc":
        return pa.timestamp("us", tz="UTC")
    return getattr(pa, name)()


def _fix_null_typed_columns(table: Any, column_oids: list[int]) -> Any:
    """Replace any all-``null``-inferred column with its real Postgres type."""
    import pyarrow as pa

    schema = table.schema
    changed = False
    for i, field in enumerate(schema):
        if pa.types.is_null(field.type):
            schema = schema.set(i, field.with_type(_arrow_type_for_oid(column_oids[i])))
            changed = True
    return table.cast(schema) if changed else table


def export_to_csv(
    conn: psycopg.Connection,
    rel: ExportRelation,
    out_path: Path,
    season: int | None = None,
) -> int:
    """Stream relation rows directly to a CSV file."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    query, params = _build_select_query(rel, season=season)
    _set_read_only_repeatable_read(conn)

    row_count = 0
    try:
        with conn.cursor(name=f"export_csv_{rel.schema}_{rel.table}") as cur:
            cur.itersize = 5000
            cur.execute(query, params)
            headers = [desc.name for desc in cur.description or []]
            with out_path.open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                while True:
                    rows = cur.fetchmany(5000)
                    if not rows:
                        break
                    writer.writerows(rows)
                    row_count += len(rows)
    finally:
        conn.commit()
    return row_count


def export_to_parquet(
    conn: psycopg.Connection,
    rel: ExportRelation,
    out_path: Path,
    season: int | None = None,
    order_by: tuple[str, ...] | None = None,
) -> int:
    """Stream relation rows to an Apache Parquet file using PyArrow.

    ``order_by`` (e.g. ``rel.primary_key``) makes row order deterministic
    across repeated exports of the same database state.
    """
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ImportError:
        raise RuntimeError(
            "Parquet export requires pyarrow. "
            "Install with `pip install 'mlb-baseball[export]'` or `uv sync --extra export`."
        ) from None

    out_path.parent.mkdir(parents=True, exist_ok=True)
    query, params = _build_select_query(rel, season=season, order_by=order_by)
    _set_read_only_repeatable_read(conn)

    row_count = 0
    try:
        with conn.cursor(name=f"export_parquet_{rel.schema}_{rel.table}") as cur:
            cur.itersize = 5000
            cur.execute(query, params)
            headers = [desc.name for desc in cur.description or []]
            column_oids = [desc.type_code for desc in cur.description or []]
            # PostgreSQL numeric (OID 1700) arrives as Python Decimal, and
            # pyarrow infers a per-batch-minimal decimal128(precision, scale)
            # from whatever values happen to be in that batch. Two batches of
            # the same column can infer *different*, mutually-incompatible
            # precisions -- confirmed for real against production
            # gold.pitching_season (ArrowInvalid: "Decimal value does not fit
            # in precision 22"), not just a null-batch edge case. Converting
            # to float before pyarrow ever sees the value makes every batch
            # infer the same stable float64, matching the float64 fallback
            # already used for an all-null numeric batch (see
            # _PG_OID_TO_ARROW_TYPE's SHORTCUT note: this project's numeric
            # columns are unconstrained-precision rate stats, ample headroom).
            _NUMERIC_OID = 1700
            numeric_col_indexes = [i for i, oid in enumerate(column_oids) if oid == _NUMERIC_OID]
            writer: pq.ParquetWriter | None = None
            try:
                while True:
                    rows = cur.fetchmany(5000)
                    if not rows:
                        if writer is None:
                            # Empty result set -> write empty table with column names
                            empty_pydict: dict[str, list[Any]] = {h: [] for h in headers}
                            table = pa.Table.from_pydict(empty_pydict)
                            table = _fix_null_typed_columns(table, column_oids)
                            pq.write_table(table, out_path)
                        break
                    if numeric_col_indexes:
                        rows = [
                            tuple(
                                float(value)
                                if i in numeric_col_indexes and value is not None
                                else value
                                for i, value in enumerate(row)
                            )
                            for row in rows
                        ]
                    # Transpose rows to column arrays for pyarrow.Table
                    cols = list(zip(*rows, strict=True))
                    batch_dict = {
                        h: list(col_data) for h, col_data in zip(headers, cols, strict=True)
                    }
                    table = pa.Table.from_pydict(batch_dict)
                    table = _fix_null_typed_columns(table, column_oids)
                    if writer is None:
                        writer = pq.ParquetWriter(out_path, table.schema)
                    elif table.schema != writer.schema:
                        table = table.cast(writer.schema)
                    writer.write_table(table)
                    row_count += len(rows)
            finally:
                if writer is not None:
                    writer.close()
    finally:
        conn.commit()
    return row_count


def export_to_xlsx(
    conn: psycopg.Connection,
    rel: ExportRelation,
    out_path: Path,
    season: int | None = None,
) -> int:
    """Stream relation rows to an Excel (.xlsx) workbook using openpyxl."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError(
            "Excel export requires openpyxl. "
            "Install with `pip install 'mlb-baseball[export]'` or `uv sync --extra export`."
        ) from None

    # Isolation must be set before the first query of the transaction, so do it
    # up front -- the row-count pre-check below then runs inside the same
    # read-only, repeatable-read snapshot as the streaming export.
    _set_read_only_repeatable_read(conn)

    # Pre-check row count against Excel sheet capacity (1,048,576 rows)
    count_query, count_params = _build_count_query(rel, season=season)
    with conn.cursor() as cur:
        cur.execute(count_query, count_params)
        (total_rows,) = fetch_one(cur)

    if total_rows > MAX_EXCEL_DATA_ROWS:
        raise ValueError(
            f"Excel export exceeds maximum sheet limit (1,048,576 rows; "
            f"relation query returned {total_rows:,} rows). "
            "Use .parquet or .csv instead."
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title=rel.table[:31])

    query, params = _build_select_query(rel, season=season)

    row_count = 0
    try:
        with conn.cursor(name=f"export_xlsx_{rel.schema}_{rel.table}") as cur:
            cur.itersize = 5000
            cur.execute(query, params)
            headers = [desc.name for desc in cur.description or []]
            ws.append(headers)
            while True:
                rows = cur.fetchmany(5000)
                if not rows:
                    break
                for r in rows:
                    formatted_row = []
                    for val in r:
                        if hasattr(val, "tzinfo") and val.tzinfo is not None:
                            formatted_row.append(val.isoformat())
                        else:
                            formatted_row.append(val)
                    ws.append(formatted_row)
                row_count += len(rows)
        wb.save(out_path)
    finally:
        conn.commit()
    return row_count


def export_relation(
    conn: psycopg.Connection,
    relation: str,
    *,
    format: str | None = None,
    out_path: Path | str | None = None,
    season: int | None = None,
) -> tuple[Path, int]:
    """Export a single allow-listed relation to the target format and file path.

    Returns a tuple of (resolved_out_path, row_count).
    """
    rel = resolve_relation(relation)

    # Infer or default format and file path
    fmt = format.lower() if format else None
    if out_path is not None:
        target_path = Path(out_path)
        if fmt is None:
            ext = target_path.suffix.lstrip(".").lower()
            if ext in ("csv", "xlsx", "parquet"):
                fmt = ext
            else:
                fmt = "parquet"
    else:
        fmt = fmt or "parquet"
        target_path = Path(f"{rel.schema}.{rel.table}.{fmt}")

    if fmt == "csv":
        count = export_to_csv(conn, rel, target_path, season=season)
    elif fmt == "xlsx":
        count = export_to_xlsx(conn, rel, target_path, season=season)
    elif fmt == "parquet":
        count = export_to_parquet(conn, rel, target_path, season=season)
    else:
        raise ValueError(
            f"Unsupported format {format!r}. Supported formats: 'csv', 'xlsx', 'parquet'."
        )

    return target_path, count


def _relation_exists(conn: psycopg.Connection, schema: str, table: str) -> bool:
    """Check if a table or view exists in the database."""
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT 1 FROM information_schema.tables WHERE table_schema = %s AND table_name = %s
            UNION ALL
            SELECT 1 FROM information_schema.views WHERE table_schema = %s AND table_name = %s
            """,
            (schema, table, schema, table),
        )
        return cur.fetchone() is not None


def export_bundle(
    conn: psycopg.Connection,
    *,
    profile: str = "public_safe",
    out_dir: Path | str = Path("export_bundle"),
    make_zip: bool = False,
) -> Path:
    """Export a rights-filtered collection of relations into a directory bundle."""
    if profile != "public_safe":
        raise ValueError(
            f"Unsupported export bundle profile {profile!r}. "
            "Currently only 'public_safe' is supported."
        )

    bundle_dir = Path(out_dir)
    bundle_dir.mkdir(parents=True, exist_ok=True)

    manifest_relations = []
    for rel in RELATIONS:
        if rel.profile != profile:
            continue

        if not _relation_exists(conn, rel.schema, rel.table):
            logger.warning(
                "Skipping absent relation %s from %s bundle", rel.qualified_name, profile
            )
            continue

        target_file = bundle_dir / f"{rel.schema}.{rel.table}.parquet"
        row_count = export_to_parquet(conn, rel, target_file)
        manifest_relations.append(
            {
                "relation": rel.qualified_name,
                "file": target_file.name,
                "row_count": row_count,
                "source_rights": rel.rights_note,
            }
        )

    # Write MANIFEST.json
    manifest_data = {
        "generated_at": datetime.now(UTC).isoformat(),
        "profile": profile,
        "relations": manifest_relations,
        "attribution": RETROSHEET_ATTRIBUTION,
    }
    manifest_path = bundle_dir / "MANIFEST.json"
    manifest_path.write_text(json.dumps(manifest_data, indent=2) + "\n", encoding="utf-8")

    if make_zip:
        zip_path = (
            bundle_dir.with_suffix(".zip") if not str(bundle_dir).endswith(".zip") else bundle_dir
        )
        if zip_path == bundle_dir:
            zip_path = bundle_dir.parent / f"{bundle_dir.name}.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for file_path in bundle_dir.rglob("*"):
                if file_path.is_file():
                    zf.write(file_path, arcname=file_path.relative_to(bundle_dir))
        return zip_path

    return bundle_dir


_DATASET_CARD_TEMPLATE = """\
# MLB Research Statistic Backbone

Grain-complete batting and pitching box-score, season, team-season, and \
career statistics, computed entirely from Retrosheet play-by-play data.

## Source

Every value in this dataset is derived from `raw.retrosheet_event` \
(Retrosheet play-by-play, 1910-2025). {attribution}

## Coverage

- Seasons: {season_min}-{season_max} (regular season only)
- Tables: {table_list}
- Excluded from this release (see `manifest.json`'s `excluded` list for the \
recorded reason): {excluded_list}

## Licence

Retrosheet's data use policy permits giving away, selling, and commercial \
products built on its data, with attribution. This repository's own code is \
AGPL-3.0-or-later. See `docs/SOURCE_RIGHTS.md` in the source repository for \
the full source-rights ledger.

## Schema version

`{schema_version}`. See `manifest.json` for the exact column names and types \
per table.

## Links

- Source repository: {repo_url}
- Methodology / honest limitations: {repo_url}/blob/main/docs/RESEARCH.md
"""


def _write_dataset_card(
    out_dir: Path,
    manifest_data: dict[str, Any],
    *,
    season_bounds: tuple[int | None, int | None],
    repo_url: str = "https://github.com/cbwinslow/mlb-baseball",
) -> Path:
    """Write the HF-facing dataset card (README.md) for a backbone bundle."""
    season_min, season_max = season_bounds
    card = _DATASET_CARD_TEMPLATE.format(
        attribution=RETROSHEET_ATTRIBUTION,
        season_min=season_min if season_min is not None else "unknown",
        season_max=season_max if season_max is not None else "unknown",
        table_list=", ".join(t["table"] for t in manifest_data["tables"]),
        excluded_list=(
            ", ".join(e["table"] for e in manifest_data["excluded"])
            if manifest_data["excluded"]
            else "none"
        ),
        schema_version=manifest_data["schema_version"],
        repo_url=repo_url,
    )
    card_path = out_dir / "README.md"
    card_path.write_text(card, encoding="utf-8")
    return card_path


def export_backbone_bundle(
    conn: psycopg.Connection,
    *,
    out_dir: Path | str = Path("backbone_bundle"),
) -> Path:
    """Export the publishable subset of the grain-complete statistic backbone.

    Writes ``<out_dir>/data/<table>.parquet`` per eligible table (deterministic
    row order via each table's primary key), ``<out_dir>/manifest.json``, and
    ``<out_dir>/README.md`` (the dataset card) -- the layout Hugging Face
    Datasets expects at a repo root. Ineligible candidates (see
    ``BACKBONE_EXCLUDED``, reviewed in
    ``openspec/changes/delivery-surface/rights-review.md``) are recorded in
    the manifest's ``excluded`` list with their reason, not written.

    SHORTCUT: each table (and the season-bounds query) runs in its own
    read-only repeatable-read transaction via ``export_to_parquet``, the same
    pattern ``export_bundle`` already uses -- not one shared snapshot across
    the whole bundle. A concurrent write (e.g. `mlb report` re-running)
    during export could make different tables reflect different points in
    time. Revisit if `export_to_parquet` gains a shared-transaction mode
    usable by every caller (single-relation CLI export, `export_bundle`,
    and this function), or if a real cross-table inconsistency is observed.
    """
    import pyarrow.parquet as pq

    bundle_dir = Path(out_dir)
    data_dir = bundle_dir / "data"
    # A stale file from a prior run (e.g. player_season.parquet written before
    # this exclusion gate existed) would otherwise survive into a re-run and
    # get published anyway -- HfApi().upload_folder() uploads the whole
    # directory regardless of what the fresh manifest says is excluded.
    if data_dir.exists():
        shutil.rmtree(data_dir)
    data_dir.mkdir(parents=True)

    manifest_tables: list[dict[str, Any]] = []
    manifest_excluded: list[dict[str, Any]] = []

    for qualified_name in BACKBONE_CANDIDATES:
        bare_table = qualified_name.split(".", 1)[1]

        if qualified_name in BACKBONE_EXCLUDED:
            reason = BACKBONE_EXCLUDED[qualified_name]
            logger.info("Excluding %s from backbone bundle: %s", qualified_name, reason)
            manifest_excluded.append(
                {"table": bare_table, "relation": qualified_name, "reason": reason}
            )
            continue

        rel = resolve_relation(qualified_name)
        if not _relation_exists(conn, rel.schema, rel.table):
            raise RuntimeError(
                f"Backbone export: required table {qualified_name!r} is missing from "
                "the database. Run migrations/build the statistic backbone first."
            )

        target_file = data_dir / f"{bare_table}.parquet"
        row_count = export_to_parquet(conn, rel, target_file, order_by=rel.primary_key)
        schema = pq.read_schema(target_file)
        manifest_tables.append(
            {
                "table": bare_table,
                "relation": qualified_name,
                "file": f"data/{bare_table}.parquet",
                "row_count": row_count,
                "columns": [
                    {"name": name, "type": str(schema.field(name).type)} for name in schema.names
                ],
            }
        )

    season_bounds: tuple[int | None, int | None] = (None, None)
    with conn.cursor() as cur:
        cur.execute("SELECT min(season), max(season) FROM gold.batting_game")
        row = cur.fetchone()
        if row is not None:
            season_bounds = (row[0], row[1])
    conn.commit()

    manifest_data: dict[str, Any] = {
        "schema_version": BACKBONE_SCHEMA_VERSION,
        "generated_at": datetime.now(UTC).isoformat(),
        "preset": "backbone",
        "tables": manifest_tables,
        "excluded": manifest_excluded,
        "attribution": RETROSHEET_ATTRIBUTION,
    }
    manifest_path = bundle_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest_data, indent=2, sort_keys=True) + "\n", "utf-8")

    _write_dataset_card(bundle_dir, manifest_data, season_bounds=season_bounds)

    return bundle_dir


def health_check() -> list[Check]:
    """Operational health check for the research data exporter layer."""
    checks: list[Check] = []
    if not ALLOWLIST:
        checks.append(Check("export allowlist", False, "No export relations registered"))
    else:
        public_count = sum(1 for r in RELATIONS if r.profile == "public_safe")
        checks.append(
            Check(
                "export allowlist",
                True,
                f"{len(ALLOWLIST)} relations registered ({public_count} public_safe)",
            )
        )

    try:
        import pyarrow  # noqa: F401

        checks.append(Check("export pyarrow dependency", True, "pyarrow installed"))
    except ImportError:
        checks.append(
            Check(
                "export pyarrow dependency",
                False,
                "pyarrow missing — install with `pip install 'mlb-baseball[export]'`",
            )
        )

    try:
        import openpyxl  # noqa: F401

        checks.append(Check("export openpyxl dependency", True, "openpyxl installed"))
    except ImportError:
        checks.append(
            Check(
                "export openpyxl dependency",
                False,
                "openpyxl missing — install with `pip install 'mlb-baseball[export]'`",
            )
        )

    return checks
